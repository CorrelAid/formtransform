"""Render each question type as it appears to a respondent and crop the image
to the question presentation.

Two engines:
  * limesurvey — imports each entity's ``tsv.tsv`` into the repo's live
    LimeSurvey stack (see ``codegen.limesurvey_stack``), activates it, and
    screenshots the public survey page with Playwright. An already-running stack
    is reused; otherwise one is started and torn down again.
  * xlsform    — uploads each entity's ``xlsform.xlsx`` to ODK XLSForm Online,
    renders it in Enketo, and screenshots the question. NOTE: this sends the
    form definition to the public ODK/Enketo staging services.

Heavy (Docker + a headless browser + network), so it is gated behind
``codegen --screenshots`` and never runs as part of the default generation.

Requirements: docker + docker compose, node, and Playwright with Chromium
(``npm i -g playwright`` or a local install; ``playwright install chromium``).
"""

from __future__ import annotations

import csv
import json
import subprocess
import tempfile
from pathlib import Path

from . import limesurvey_stack as ls_stack

SHOT_DIR = Path(__file__).resolve().parent / "screenshots"
SHOT_LS = SHOT_DIR / "shot_limesurvey.mjs"
SHOT_XLS = SHOT_DIR / "shot_xlsform.mjs"


def _run(cmd: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, check=True, text=True, **kw)


# --- Engines ---------------------------------------------------------------
def _has_datetime_question(tsv: Path) -> bool:
    """Does this TSV contain a LimeSurvey date/time question (`type/scale` == D)?

    Read straight off the blessed TSV's own type code — the very thing that
    decides which widget LimeSurvey renders. No slug list and no registry
    lookup, so any type that maps to `D` gets the `--picker` treatment for free.
    """
    with open(tsv, newline="") as f:
        for row in csv.DictReader(f, delimiter="\t"):
            if (row.get("class") or "").strip() == "Q" and (row.get("type/scale") or "").strip() == "D":
                return True
    return False


def _shot_limesurvey(entities: list[tuple[str, Path]], out_dir: Path) -> dict[str, str]:
    report: dict[str, str] = {}
    started = ls_stack.ensure_up()
    try:
        ls = ls_stack.client()
        for slug, edir in entities:
            tsv = edir / "tsv.tsv"
            if not tsv.exists():
                continue
            try:
                survey_id = ls_stack.import_and_activate(ls, tsv, f"{slug}_demo")
                out = out_dir / slug / "limesurvey.png"
                out.parent.mkdir(parents=True, exist_ok=True)
                cmd = ["node", str(SHOT_LS), ls_stack.public_url(survey_id), str(out)]
                if slug.endswith("_other"):
                    cmd.append("--other")
                if _has_datetime_question(tsv):
                    cmd.append("--picker")
                _run(cmd)
                report[slug] = "ok"
            except Exception as e:
                report[slug] = f"error: {e}"
    finally:
        if started:
            ls_stack.compose_down()
    return report


_FROM_FILE_TYPES = {"select_one_from_file": "select_one", "select_multiple_from_file": "select_multiple"}


def _render_xlsx(edir: Path, tmp_dir: Path) -> Path:
    """Return the xlsx to upload for rendering.

    ``select_*_from_file`` questions reference an external CSV that Enketo's
    preview cannot resolve (no media is attached), so it errors. For the
    screenshot we inline the vocabulary as regular ``select_one``/
    ``select_multiple`` choices — the presentation is identical to a deployed
    from-file question once its media is present. Every other type uses the
    committed ``xlsform.xlsx`` unchanged.
    """
    committed = edir / "xlsform.xlsx"
    src = edir / "xlsform.json"
    if not src.exists():
        return committed
    xlsform = json.loads(src.read_text())
    survey = xlsform.get("survey", [])
    needs_inline = any(str(r.get("type", "")).split()[0] in _FROM_FILE_TYPES for r in survey)
    if not needs_inline:
        return committed

    from openpyxl import Workbook

    vocab_dir = edir.parent.parent / "vocab"
    new_survey: list[dict] = []
    choices: list[dict] = list(xlsform.get("choices", []))
    for r in survey:
        parts = str(r.get("type", "")).split()
        base = parts[0] if parts else ""
        if base in _FROM_FILE_TYPES and len(parts) > 1:
            list_name = Path(parts[1]).stem
            nr = dict(r)
            nr["type"] = f"{_FROM_FILE_TYPES[base]} {list_name}"
            new_survey.append(nr)
            with open(vocab_dir / parts[1], newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    choices.append(
                        {
                            "list_name": list_name,
                            "name": row.get("code") or row.get("name"),
                            "label": row.get("label"),
                        }
                    )
        else:
            new_survey.append(r)

    wb = Workbook()
    ws_s = wb.active
    ws_s.title = "survey"
    s_cols = sorted({k for row in new_survey for k in row}) or ["type", "name", "label"]
    ws_s.append(s_cols)
    for row in new_survey:
        ws_s.append([row.get(c, "") for c in s_cols])
    ws_c = wb.create_sheet("choices")
    c_cols = sorted({k for row in choices for k in row}) or ["list_name", "name", "label"]
    ws_c.append(c_cols)
    for row in choices:
        ws_c.append([row.get(c, "") for c in c_cols])
    ws_set = wb.create_sheet("settings")
    ws_set.append(["form_title", "form_id", "default_language"])
    ws_set.append([edir.name, edir.name, "default"])

    dest = tmp_dir / f"{edir.name}.xlsx"
    wb.save(dest)
    return dest


def _shot_xlsform(entities: list[tuple[str, Path]], out_dir: Path) -> dict[str, str]:
    report: dict[str, str] = {}
    tmp_dir = Path(tempfile.mkdtemp(prefix="cdl-shots-xlsx-"))
    for slug, edir in entities:
        if not (edir / "xlsform.xlsx").exists():
            continue
        try:
            xlsx = _render_xlsx(edir, tmp_dir)
            out = out_dir / slug / "xlsform.png"
            out.parent.mkdir(parents=True, exist_ok=True)
            cmd = ["node", str(SHOT_XLS), str(xlsx), str(out)]
            if slug.endswith("_other"):
                cmd.append("--other")
            _run(cmd)
            report[slug] = "ok"
        except Exception as e:
            report[slug] = f"error: {e}"
    return report


def generate_screenshots(
    base_dir: Path,
    out_dir: Path,
    engines: tuple[str, ...] = ("limesurvey", "xlsform"),
    only: list[str] | None = None,
) -> dict[str, dict[str, str]]:
    """Render screenshots for every entity on disk (optionally filtered by ``only``).

    Writes ``<out_dir>/<slug>/{limesurvey,xlsform}.png``.
    """
    entities = sorted((d.name, d) for d in (base_dir / "registry" / "entities").iterdir() if d.is_dir())
    if only:
        wanted = set(only)
        entities = [(s, d) for s, d in entities if s in wanted]

    out: dict[str, dict[str, str]] = {}
    if "limesurvey" in engines:
        out["limesurvey"] = _shot_limesurvey(entities, out_dir)
    if "xlsform" in engines:
        out["xlsform"] = _shot_xlsform(entities, out_dir)
    return out
